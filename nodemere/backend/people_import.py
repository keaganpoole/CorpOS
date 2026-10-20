"""Parsing and validation helpers for append-only People imports."""

import csv
import hashlib
import io
import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from fastapi import HTTPException
from openpyxl import load_workbook


MAX_IMPORT_BYTES = 10 * 1024 * 1024
MAX_IMPORT_ROWS = 20000
MAX_IMPORT_COLUMNS = 150
IMPORT_MARKER = "__nodemere_import_key"

# Server owned allowlist. Operational IDs, Stripe IDs, and activity timestamps
# cannot be supplied by a spreadsheet.
CORE_FIELDS = {
    "full_name": ("Full Name", "text"),
    "first_name": ("First Name", "text"),
    "last_name": ("Last Name", "text"),
    "phone": ("Phone", "text"),
    "email": ("Email", "text"),
    "street_address": ("Street Address", "text"),
    "city": ("City", "text"),
    "state": ("State", "text"),
    "zip_code": ("Zip Code", "text"),
    "preferred_contact_method": ("Preferred Contact Method", "text"),
    "preferred_language": ("Preferred Language", "text"),
    "best_time_to_contact": ("Best Time To Contact", "text"),
    "source": ("Source", "text"),
    "lead_source_detail": ("Source Detail", "text"),
    "status": ("Status", "text"),
    "notes": ("Notes", "text"),
    "special_instructions": ("Special Instructions", "text"),
    "tags": ("Tags", "multi_select"),
    "do_not_call": ("Do Not Call", "boolean"),
    "do_not_text": ("Do Not Text", "boolean"),
    "consent_call": ("Documented AI Call Consent", "boolean"),
    "consent_call_source": ("AI Call Consent Source", "text"),
    "consent_call_recorded_at": ("AI Call Consent Timestamp", "timestamp"),
    "consent_call_scope": ("AI Call Consent Scope", "text"),
    "consent_sms": ("Documented SMS Consent", "boolean"),
    "consent_sms_source": ("SMS Consent Source", "text"),
    "consent_sms_recorded_at": ("SMS Consent Timestamp", "timestamp"),
    "consent_sms_scope": ("SMS Consent Scope", "text"),
}


def _scalar(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, list) and all(isinstance(part, (str, int, float, bool)) for part in value):
        return value
    raise HTTPException(400, "Import cells must contain text, numbers, dates, or simple lists.")


def _finish_rows(headers, rows):
    headers = [str(header or "").strip() for header in headers]
    if not headers or any(not header for header in headers):
        raise HTTPException(400, "Every import column needs a header.")
    if len(headers) > MAX_IMPORT_COLUMNS:
        raise HTTPException(400, "This file has too many columns.")
    if len(set(headers)) != len(headers):
        raise HTTPException(400, "Column headers must be unique.")
    for row in rows:
        row["values"] = {str(key).strip(): value for key, value in row["values"].items()}
    if not rows:
        raise HTTPException(400, "The file has no records to import.")
    if len(rows) > MAX_IMPORT_ROWS:
        raise HTTPException(400, f"Import up to {MAX_IMPORT_ROWS:,} records at a time.")
    return {"headers": headers, "rows": rows}


def parse_people_import(content: bytes, filename: str):
    if not content or len(content) > MAX_IMPORT_BYTES:
        raise HTTPException(400, "Choose a file under 10 MB.")
    suffix = Path(filename or "").suffix.lower()
    rows = []
    try:
        if suffix == ".csv":
            stream = io.StringIO(content.decode("utf-8-sig"), newline="")
            sample = stream.read(4096)
            stream.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            reader = csv.reader(stream, dialect)
            headers = next(reader, [])
            for line, values in enumerate(reader, start=2):
                if not any(str(value).strip() for value in values):
                    continue
                if len(values) != len(headers):
                    raise HTTPException(400, f"Row {line} has a different number of columns.")
                rows.append({"line": line, "values": dict(zip(headers, values))})
                if len(rows) > MAX_IMPORT_ROWS:
                    break
        elif suffix == ".xlsx":
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                sheet = workbook.active
                iterator = sheet.iter_rows(values_only=True)
                headers = list(next(iterator, ()))
                for line, values in enumerate(iterator, start=2):
                    if not any(value is not None and str(value).strip() for value in values):
                        continue
                    if any(value is not None for value in values[len(headers):]):
                        raise HTTPException(400, f"Row {line} has values outside the header columns.")
                    rows.append({"line": line, "values": {
                        str(header): _scalar(value) for header, value in zip(headers, values)
                    }})
                    if len(rows) > MAX_IMPORT_ROWS:
                        break
            finally:
                workbook.close()
        elif suffix == ".json":
            parsed = json.loads(content.decode("utf-8-sig"))
            if isinstance(parsed, dict):
                parsed = parsed.get("people", parsed.get("records"))
            if not isinstance(parsed, list):
                raise HTTPException(400, "JSON must be an array of people, or an object with a people or records array.")
            headers = []
            for line, record in enumerate(parsed, start=1):
                if not isinstance(record, dict):
                    raise HTTPException(400, f"JSON record {line} must be an object.")
                values = {key: value for key, value in record.items() if key != "custom_fields"}
                custom = record.get("custom_fields")
                if custom is not None:
                    if not isinstance(custom, dict):
                        raise HTTPException(400, f"JSON record {line} has invalid custom_fields.")
                    values.update({key: value for key, value in custom.items()
                                   if isinstance(key, str) and key.startswith("custom_") and key not in values})
                for key in values:
                    if key not in headers:
                        headers.append(key)
                rows.append({"line": line, "values": {key: _scalar(value) for key, value in values.items()}})
                if len(rows) > MAX_IMPORT_ROWS or len(headers) > MAX_IMPORT_COLUMNS:
                    break
        else:
            raise HTTPException(400, "Choose a CSV, XLSX, or JSON file.")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(400, "The file could not be read. Check its format and try again.") from exc
    result = _finish_rows(headers, rows)
    result["file_hash"] = hashlib.sha256(content).hexdigest()
    return result


def import_fields(custom_schema):
    fields = [{"key": key, "label": label, "type": kind} for key, (label, kind) in CORE_FIELDS.items()]
    fields.extend({"key": row["field_key"], "label": row.get("label") or row["field_key"], "type": row.get("field_type") or "text"}
                  for row in custom_schema if isinstance(row.get("field_key"), str) and row["field_key"].startswith("custom_") and row.get("field_type") != "docs")
    return fields


def _boolean(value):
    if isinstance(value, bool):
        return value
    lowered = str(value).strip().lower()
    if lowered in {"true", "yes", "1", "y"}:
        return True
    if lowered in {"false", "no", "0", "n"}:
        return False
    raise ValueError("Use Yes or No for yes/no fields")


def normalize_import_record(values, field_types):
    if not isinstance(values, dict):
        raise ValueError("This row is not a record")
    payload = {}
    for key, value in values.items():
        if key not in field_types:
            raise ValueError("An import field is no longer available")
        if value is None or value == "":
            continue
        kind = field_types[key]
        if kind == "boolean":
            payload[key] = _boolean(value)
        elif kind == "number":
            try:
                number = Decimal(str(value).replace(",", ""))
                if not number.is_finite():
                    raise InvalidOperation()
                payload[key] = float(number)
            except (InvalidOperation, ValueError):
                raise ValueError("A number is invalid") from None
        elif kind == "multi_select":
            parts = value if isinstance(value, list) else re.split(r"[,;]", str(value))
            payload[key] = [str(part).strip() for part in parts if str(part).strip()]
        elif kind in {"date", "timestamp"}:
            text = str(value).strip()
            try:
                parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
            except ValueError:
                try:
                    parsed = date.fromisoformat(text)
                except ValueError:
                    raise ValueError("A date or timestamp is invalid") from None
            payload[key] = parsed.date().isoformat() if kind == "date" and isinstance(parsed, datetime) else parsed.isoformat()
        else:
            if isinstance(value, (dict, list)):
                raise ValueError("A cell has an unsupported value")
            text = str(value).strip()
            if len(text) > (20000 if key in {"notes", "special_instructions"} else 500):
                raise ValueError("A value is too long")
            payload[key] = text
    full_name = payload.pop("full_name", "")
    if full_name and not (payload.get("first_name") or payload.get("last_name")):
        first, _, last = full_name.partition(" ")
        payload["first_name"], payload["last_name"] = first, last
    if not any(payload.get(key) for key in ("first_name", "last_name", "phone", "email")):
        raise ValueError("Add a name, phone, or email")
    if payload.get("email") and not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", payload["email"]):
        raise ValueError("The email address is invalid")
    for consent in ("call", "sms"):
        if payload.get(f"consent_{consent}") is True and not all(payload.get(f"consent_{consent}_{part}") for part in ("source", "recorded_at", "scope")):
            raise ValueError("Documented consent needs a source, timestamp, and scope")
    return payload


def identity_keys(person):
    keys = []
    email = str(person.get("email") or "").strip().lower()
    if email:
        keys.append("email:" + email)
    digits = re.sub(r"\D", "", str(person.get("phone") or ""))
    if len(digits) == 11 and digits.startswith("1"):
        digits = digits[1:]
    if len(digits) >= 7:
        keys.append("phone:" + digits)
    return keys


def row_import_key(file_hash, line):
    if not re.fullmatch(r"[0-9a-f]{64}", str(file_hash)) or not isinstance(line, int) or line < 1:
        raise ValueError("Import identity is invalid")
    return hashlib.sha256(f"{file_hash}:{line}".encode()).hexdigest()


def prepare_people_import(rows, file_hash, field_types, owner, business_id, known_keys, known_markers):
    if not isinstance(rows, list) or len(rows) > MAX_IMPORT_ROWS:
        raise HTTPException(400, "Import up to 20,000 records at a time.")
    if not isinstance(file_hash, str) or not re.fullmatch(r"[0-9a-f]{64}", file_hash):
        raise HTTPException(400, "Import file identity is invalid.")
    seen_keys, seen_markers = set(known_keys), set(known_markers)
    decisions, prepared = [], []
    for item in rows:
        line = item.get("line") if isinstance(item, dict) else None
        try:
            marker = row_import_key(file_hash, line)
            values = normalize_import_record(item.get("values"), field_types)
            keys = identity_keys(values)
        except (ValueError, TypeError) as exc:
            decisions.append({"line": line, "status": "invalid", "reason": str(exc)})
            continue
        if marker in seen_markers or any(key in seen_keys for key in keys):
            decisions.append({"line": line, "status": "duplicate"})
            continue
        seen_markers.add(marker)
        seen_keys.update(keys)
        custom = {key: values.pop(key) for key in list(values) if key.startswith("custom_")}
        custom[IMPORT_MARKER] = marker
        values.update({"user_id": owner, "business_id": business_id, "custom_fields": custom})
        prepared.append(values)
        decisions.append({"line": line, "status": "ready"})
    return decisions, prepared
